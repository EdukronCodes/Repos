from __future__ import annotations

import json
import mimetypes
import subprocess
import zipfile
from pathlib import Path
from typing import Any

import pandas as pd


def json_safe(value: Any) -> Any:
    if value is None:
        return None
    if hasattr(value, "item"):
        value = value.item()
    try:
        return None if pd.isna(value) else value
    except (TypeError, ValueError):
        return value


class FileParser:
    """Convert supported uploads into shared text and tabular records."""

    SPREADSHEETS = {".xlsx", ".xls", ".xlsm", ".ods"}
    TEXT = {".txt", ".md", ".log", ".json", ".xml", ".yaml", ".yml", ".html", ".htm", ".csv"}
    IMAGES = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".webp"}
    AUDIO_VIDEO = {".mp3", ".wav", ".m4a", ".mp4", ".mov", ".avi", ".mkv"}
    ARCHIVES = {".zip", ".tar", ".gz", ".rar", ".7z"}

    def parse(self, path: Path) -> tuple[list[dict[str, Any]], dict[str, list[dict[str, Any]]]]:
        suffix = path.suffix.lower()
        if suffix in self.SPREADSHEETS:
            workbook = pd.read_excel(path, sheet_name=None)
            return self._tables(workbook)
        if suffix == ".csv":
            frame = pd.read_csv(path)
            return self._tables({path.stem: frame})
        if suffix in self.TEXT:
            text = path.read_text(encoding="utf-8", errors="replace")
            return [{"name": path.name, "kind": "text", "text": text}], {}
        if suffix == ".pdf":
            return [{"name": path.name, "kind": "pdf", "text": self._pdf(path)}], {}
        if suffix == ".docx":
            return [{"name": path.name, "kind": "document", "text": self._docx(path)}], {}
        if suffix in {".pptx", ".ppt"}:
            return [{"name": path.name, "kind": "presentation", "text": self._pptx(path)}], {}
        if suffix in self.IMAGES:
            return [{"name": path.name, "kind": "image", "text": self._image(path)}], {}
        if suffix in self.AUDIO_VIDEO:
            return [{"name": path.name, "kind": "media", "text": self._media(path)}], {}
        if suffix in self.ARCHIVES:
            return [{"name": path.name, "kind": "archive", "text": self._archive(path)}], {}
        return [{"name": path.name, "kind": "binary", "text": f"File type: {mimetypes.guess_type(path.name)[0] or 'unknown'}"}], {}

    def _tables(self, workbook: dict[str, pd.DataFrame]) -> tuple[list[dict[str, Any]], dict[str, list[dict[str, Any]]]]:
        documents = [{"name": name, "kind": "table", "text": frame.to_csv(index=False)} for name, frame in workbook.items()]
        return documents, {name: self._records(frame) for name, frame in workbook.items()}

    @staticmethod
    def _records(frame: pd.DataFrame) -> list[dict[str, Any]]:
        return [{str(key): json_safe(value) for key, value in row.items()} for row in frame.to_dict(orient="records")]

    @staticmethod
    def _pdf(path: Path) -> str:
        try:
            from pypdf import PdfReader
            return "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
        except (ImportError, Exception) as error:
            return f"PDF text extraction unavailable: {error}"

    @staticmethod
    def _docx(path: Path) -> str:
        try:
            from docx import Document
            return "\n".join(paragraph.text for paragraph in Document(str(path)).paragraphs)
        except ImportError:
            return "DOCX text extraction unavailable; install python-docx."

    @staticmethod
    def _pptx(path: Path) -> str:
        try:
            from pptx import Presentation
            return "\n".join(shape.text for slide in Presentation(str(path)).slides for shape in slide.shapes if hasattr(shape, "text"))
        except ImportError:
            return "PPTX text extraction unavailable; install python-pptx."

    @staticmethod
    def _image(path: Path) -> str:
        try:
            from PIL import Image
            image = Image.open(path)
            try:
                import pytesseract
                text = pytesseract.image_to_string(image).strip()
            except (ImportError, Exception):
                text = "OCR unavailable"
            return f"Image: {image.width}x{image.height}; {text}".strip()
        except ImportError:
            return "Image metadata unavailable; install pillow."

    @staticmethod
    def _media(path: Path) -> str:
        try:
            result = subprocess.run(["ffprobe", "-v", "error", "-show_format", "-show_streams", str(path)], capture_output=True, text=True, check=False)
            return f"Media metadata for {path.name}:\n{result.stdout or result.stderr}"
        except FileNotFoundError:
            return "Media metadata unavailable; install ffmpeg."

    @staticmethod
    def _archive(path: Path) -> str:
        if path.suffix.lower() == ".zip":
            with zipfile.ZipFile(path) as archive:
                return "Archive members:\n" + "\n".join(archive.namelist())
        return f"Archive inspection is available for ZIP files; received {path.suffix}."
