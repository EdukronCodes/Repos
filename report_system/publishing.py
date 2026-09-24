from __future__ import annotations

import json
from html import escape
from pathlib import Path
from typing import Any

import pandas as pd

SUPPORTED_FORMATS = ["pdf", "docx", "xlsx", "csv", "pptx", "html", "md", "json"]
MIME_TYPES = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "csv": "text/csv",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "html": "text/html",
    "md": "text/markdown",
    "json": "application/json",
}


def publish_report(report: dict[str, Any], output: Path, format_name: str) -> Path:
    format_name = format_name.lower().lstrip(".")
    if format_name not in SUPPORTED_FORMATS:
        raise ValueError(f"Unsupported format: {format_name}")
    output = output.with_suffix("." + format_name)
    output.parent.mkdir(parents=True, exist_ok=True)
    if format_name == "json":
        output.write_text(json.dumps(report, indent=2, default=str), encoding="utf-8")
    elif format_name == "md":
        output.write_text(_markdown(report), encoding="utf-8")
    elif format_name == "html":
        output.write_text(_html(report), encoding="utf-8")
    elif format_name == "csv":
        rows = [{"table": sheet["name"], **item} for sheet in report["analysis"]["tables"] for item in sheet["numeric"]]
        pd.DataFrame(rows).to_csv(output, index=False)
    elif format_name == "xlsx":
        _xlsx(report, output)
    elif format_name == "pdf":
        _pdf(report, output)
    elif format_name == "docx":
        _docx(report, output)
    elif format_name == "pptx":
        _pptx(report, output)
    return output


def report_bytes(report: dict[str, Any], format_name: str) -> tuple[bytes, str, str]:
    import tempfile
    with tempfile.TemporaryDirectory() as directory:
        path = publish_report(report, Path(directory) / "report", format_name)
        return path.read_bytes(), path.name, MIME_TYPES[format_name]


def _markdown(report: dict[str, Any]) -> str:
    sections = [f"# {report['title']}", report["goal"], *[f"- {item}" for item in report["overview"]]]
    sections.extend(f"## {sheet['name']}\n\nRows: {sheet['rows']}\n\nMissing cells: {sheet['missing_cells']}\n\nDuplicate rows: {sheet['duplicate_rows']}" for sheet in report["analysis"]["tables"])
    return "\n\n".join(sections)


def _html(report: dict[str, Any]) -> str:
    tables = "".join(f"<h2>{escape(sheet['name'])}</h2><p>Rows: {sheet['rows']} | Missing: {sheet['missing_cells']} | Duplicates: {sheet['duplicate_rows']}</p><pre>{escape(json.dumps(sheet, indent=2, default=str))}</pre>" for sheet in report["analysis"]["tables"])
    return f"<html><body><h1>{escape(report['title'])}</h1><p>{escape(report['goal'])}</p><ul>{''.join(f'<li>{escape(item)}</li>' for item in report['overview'])}</ul>{tables}</body></html>"


def _xlsx(report: dict[str, Any], output: Path) -> None:
    with pd.ExcelWriter(output) as writer:
        pd.DataFrame(report["overview"], columns=["overview"]).to_excel(writer, sheet_name="Overview", index=False)
        for sheet in report["analysis"]["tables"]:
            pd.DataFrame(sheet["numeric"]).to_excel(writer, sheet_name=str(sheet["name"])[:31] or "Data", index=False)


def _pdf(report: dict[str, Any], output: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    styles = getSampleStyleSheet()
    story = [Paragraph(report["title"], styles["Title"]), Paragraph(report["goal"], styles["BodyText"]), Spacer(1, 0.2 * inch), Paragraph(" | ".join(report["overview"]), styles["BodyText"])]
    for sheet in report["analysis"]["tables"]:
        story.extend([Spacer(1, 0.2 * inch), Paragraph(str(sheet["name"]), styles["Heading2"])])
        rows = [["Column", "Mean", "Minimum", "Maximum", "Total"]] + [[item["column"], f"{item['mean']:.2f}", f"{item['minimum']:.2f}", f"{item['maximum']:.2f}", f"{item['total']:.2f}"] for item in sheet["numeric"]]
        if len(rows) > 1:
            table = Table(rows, repeatRows=1)
            table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#19324D")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("GRID", (0, 0), (-1, -1), 0.35, colors.grey)]))
            story.append(table)
    SimpleDocTemplate(str(output), pagesize=letter).build(story)


def _docx(report: dict[str, Any], output: Path) -> None:
    from docx import Document
    document = Document()
    document.add_heading(report["title"], 0)
    document.add_paragraph(report["goal"])
    document.add_paragraph(" | ".join(report["overview"]))
    for sheet in report["analysis"]["tables"]:
        document.add_heading(str(sheet["name"]), level=1)
        document.add_paragraph(f"Rows: {sheet['rows']} | Missing cells: {sheet['missing_cells']} | Duplicates: {sheet['duplicate_rows']}")
        for item in sheet["numeric"]:
            document.add_paragraph(f"{item['column']}: mean {item['mean']:.2f}, total {item['total']:.2f}")
    document.save(output)


def _pptx(report: dict[str, Any], output: Path) -> None:
    from pptx import Presentation
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = report["title"]
    slide.placeholders[1].text = report["goal"] + "\n" + "\n".join(report["overview"])
    for sheet in report["analysis"]["tables"]:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = str(sheet["name"])
        slide.placeholders[1].text = f"Rows: {sheet['rows']}\nMissing cells: {sheet['missing_cells']}\nDuplicates: {sheet['duplicate_rows']}"
    presentation.save(output)
