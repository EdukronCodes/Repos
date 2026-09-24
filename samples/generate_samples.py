from __future__ import annotations

import shutil
import subprocess
import zipfile
from pathlib import Path

import pandas as pd

ROOT = Path(__file__).parent


def main() -> None:
    ROOT.mkdir(exist_ok=True)
    sales = pd.DataFrame({"region": ["East", "West", "East", "South"], "revenue": [12000, 18000, 15000, 9000], "orders": [42, 55, 48, 27]})
    sales.to_csv(ROOT / "sample_sales.csv", index=False)
    with pd.ExcelWriter(ROOT / "sample_workbook.xlsx") as writer:
        sales.to_excel(writer, sheet_name="Sales", index=False)
        pd.DataFrame({"status": ["Open", "Closed", "Open"], "count": [12, 25, 8]}).to_excel(writer, sheet_name="Pipeline", index=False)

    (ROOT / "sample_notes.txt").write_text("Quarterly review\nRevenue grew in the West region.\nOpen pipeline requires follow-up.", encoding="utf-8")
    (ROOT / "sample_notes.md").write_text("# Quarterly review\n\nRevenue grew in the West region.", encoding="utf-8")
    (ROOT / "sample_events.log").write_text("2026-09-01 INFO import complete\n2026-09-02 WARN missing owner", encoding="utf-8")
    (ROOT / "sample_data.json").write_text('{"team":"Sales","quarter":"Q3","target":50000}', encoding="utf-8")
    (ROOT / "sample_data.xml").write_text("<report><team>Sales</team><quarter>Q3</quarter><target>50000</target></report>", encoding="utf-8")
    (ROOT / "sample_data.yaml").write_text("team: Sales\nquarter: Q3\ntarget: 50000\n", encoding="utf-8")
    (ROOT / "sample_page.html").write_text("<html><body><h1>Quarterly Review</h1><p>Revenue grew in the West.</p></body></html>", encoding="utf-8")

    from docx import Document
    document = Document()
    document.add_heading("Quarterly Review", 0)
    document.add_paragraph("Revenue grew in the West region and the open pipeline needs follow-up.")
    document.save(ROOT / "sample_report.docx")

    from pptx import Presentation
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Quarterly Review"
    slide.placeholders[1].text = "Revenue grew in the West region."
    presentation.save(ROOT / "sample_presentation.pptx")

    from reportlab.pdfgen import canvas
    pdf = canvas.Canvas(str(ROOT / "sample_report.pdf"))
    pdf.drawString(72, 720, "Quarterly Review")
    pdf.drawString(72, 690, "Revenue grew in the West region.")
    pdf.save()

    from PIL import Image, ImageDraw
    image = Image.new("RGB", (640, 360), "white")
    ImageDraw.Draw(image).text((40, 40), "Quarterly Review - Revenue grew in the West", fill="black")
    image.save(ROOT / "sample_image.png")

    with zipfile.ZipFile(ROOT / "sample_archive.zip", "w") as archive:
        archive.write(ROOT / "sample_notes.txt", "sample_notes.txt")
        archive.write(ROOT / "sample_sales.csv", "sample_sales.csv")

    if shutil.which("ffmpeg"):
        subprocess.run(["ffmpeg", "-y", "-f", "lavfi", "-i", "sine=frequency=440:duration=1", str(ROOT / "sample_audio.wav")], capture_output=True, check=False)
        subprocess.run(["ffmpeg", "-y", "-f", "lavfi", "-i", "color=c=blue:s=320x240:d=1", str(ROOT / "sample_video.mp4")], capture_output=True, check=False)
    print(f"Created sample files in {ROOT}")


if __name__ == "__main__":
    main()
